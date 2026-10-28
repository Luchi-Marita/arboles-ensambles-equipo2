"""Build the technical report (DOCX) for the three tree ensembles.

The script is intentionally self-contained: it reads the experiment JSON
artifacts, renders equations and the cover, and writes a styled Word document.
PDF conversion is performed afterwards with Microsoft Word (see README).
"""
from __future__ import annotations

import json
import math
import re
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
FIGURES = ROOT / "artifacts" / "figures"
TABLES = ROOT / "artifacts" / "tables"
EQ = REPORTS / "equations"
REPORTS.mkdir(parents=True, exist_ok=True)
EQ.mkdir(parents=True, exist_ok=True)

REPO = "https://github.com/Luchi-Marita/arboles-ensambles-equipo2"
PAGES = "https://luchi-marita.github.io/arboles-ensambles-equipo2/"
AUTHOR = "Luchi-Marita"
TEAM = "Equipo 2"
COURSE = "Aprendizaje Supervisado · Familia 1: Árboles y Ensambles"
TODAY = date(2026, 10, 28)

NAVY = RGBColor(0x0A, 0x19, 0x28)
INK = RGBColor(0x12, 0x20, 0x2F)
MUTED = RGBColor(0x5B, 0x6B, 0x7C)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
AMBER = RGBColor(0xB4, 0x53, 0x09)
VIOLET = RGBColor(0x6D, 0x28, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_BODY = "Calibri"
FONT_DISPLAY = "Georgia"
FONT_CODE = "Consolas"


def load(path: Path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


EXTRA = load(TABLES / "extra_trees_results.json")
ADA = load(TABLES / "adaboost_results.json")
CAT = load(TABLES / "catboost_results.json")


def num(v, d=3) -> str:
    try:
        v = float(v)
    except (TypeError, ValueError):
        return "—"
    if math.isnan(v) or math.isinf(v):
        return "—"
    return f"{v:,.{d}f}".replace(",", " ")


def pct(v, d=1) -> str:
    return f"{100 * float(v):.{d}f}%"


def compact(cfg: str) -> str:
    keys = {
        "n_estimators": "n", "max_depth": "depth", "max_features": "mf",
        "min_samples_leaf": "leaf", "min_samples_split": "split",
        "bootstrap": "boot", "learning_rate": "lr", "iterations": "iter",
        "l2_leaf_reg": "l2", "random_strength": "rs",
        "bagging_temperature": "bt", "auto_class_weights": "weights",
    }
    parts = []
    for raw in str(cfg).split(" | "):
        if "=" not in raw:
            continue
        k, v = raw.split("=", 1)
        parts.append(f"{keys.get(k, k)}={v}")
    return " · ".join(parts)


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------
def _font(name: str, size: int):
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / name), size)


def build_cover() -> Path:
    W, H = 1654, 2339
    img = Image.new("RGB", (W, H), (7, 20, 33))
    px = img.load()
    for y in range(H):
        t = y / H
        r = int(7 + 10 * t)
        g = int(20 + 24 * t)
        b = int(33 + 28 * t)
        for x in range(W):
            px[x, y] = (r, g, b)

    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(overlay)
    for x in range(0, W, 96):
        d.line([(x, 0), (x, H)], fill=(255, 255, 255, 9), width=1)
    for y in range(0, H, 96):
        d.line([(0, y), (W, y)], fill=(255, 255, 255, 9), width=1)

    # decorative tree / network
    nodes = [
        (1240, 390, 46, (45, 212, 191)), (1050, 520, 34, (245, 158, 11)),
        (1430, 520, 34, (167, 139, 250)), (930, 660, 26, (45, 212, 191)),
        (1160, 680, 26, (245, 158, 11)), (1335, 660, 26, (167, 139, 250)),
        (1520, 680, 26, (45, 212, 191)), (1040, 820, 20, (45, 212, 191)),
        (1210, 830, 20, (245, 158, 11)), (1380, 820, 20, (167, 139, 250)),
        (1530, 830, 20, (45, 212, 191)),
    ]
    edges = [(0, 1), (0, 2), (1, 3), (1, 4), (2, 5), (2, 6), (3, 7), (4, 7),
             (4, 8), (5, 9), (6, 10), (6, 4)]
    for a, b in edges:
        d.line([nodes[a][:2], nodes[b][:2]], fill=(255, 255, 255, 34), width=2)
    for x, y, r, c in nodes:
        d.ellipse([x - r, y - r, x + r, y + r], fill=c + (50,), outline=c + (180,), width=2)
        d.ellipse([x - r // 2, y - r // 2, x + r // 2, y + r // 2], fill=c + (190,))

    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    d = ImageDraw.Draw(img)

    margin = 148
    d.text((margin, 120), "BOSQUE DE DECISIONES", font=_font("segoeuib.ttf", 27),
           fill=(94, 234, 212))
    d.text((margin, 162), "TRABAJO FINAL · APRENDIZAJE SUPERVISADO", font=_font("segoeui.ttf", 20),
           fill=(159, 178, 196))

    d.text((margin, 315), "Aprender", font=_font("georgiab.ttf", 150), fill=(245, 248, 250))
    d.text((margin, 470), "dividiendo", font=_font("georgiai.ttf", 150), fill=(45, 212, 191))
    d.text((margin, 625), "el espacio.", font=_font("georgiab.ttf", 150), fill=(245, 248, 250))

    sub = ("Tres ensambles, tres decisiones de negocio: Extra Trees para pronosticar\n"
           "demanda, AdaBoost para cribar diabetes y CatBoost para anticipar abandono\n"
           "universitario. Teoría, entrenamiento real y simulación en el navegador.")
    d.multiline_text((margin, 835), sub, font=_font("segoeui.ttf", 31), fill=(184, 201, 214),
                     spacing=14)

    cards = [
        ("EXTRA TREES", "RMSE 1.998", "R² 0.607 · Walmart M5", (45, 212, 191)),
        ("ADABOOST", "ROC-AUC 0.821", "F1 0.758 · CDC BRFSS", (245, 158, 11)),
        ("CATBOOST", "MACRO-F1 0.693", "AUC 0.886 · UCI Dropout", (167, 139, 250)),
    ]
    cw, gap, ch, top = 420, 31, 252, 1275
    for i, (name, metric, foot, color) in enumerate(cards):
        x0 = margin + i * (cw + gap)
        layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
        dd = ImageDraw.Draw(layer)
        dd.rounded_rectangle([x0, top, x0 + cw, top + ch], radius=22,
                             fill=(255, 255, 255, 17), outline=color + (110,), width=2)
        img = Image.alpha_composite(img.convert("RGBA"), layer).convert("RGB")
        d = ImageDraw.Draw(img)
        d.text((x0 + 28, top + 30), name, font=_font("segoeuib.ttf", 19), fill=color)
        d.text((x0 + 28, top + 72), metric, font=_font("georgiab.ttf", 40), fill=(245, 248, 250))
        d.text((x0 + 28, top + 145), foot, font=_font("segoeui.ttf", 20), fill=(174, 192, 207))
        d.line([(x0 + 28, top + 190), (x0 + cw - 28, top + 190)], fill=(255, 255, 255, 35), width=1)
        d.text((x0 + 28, top + 205), "modelo entrenado · métricas de test", font=_font("segoeui.ttf", 17),
               fill=(140, 160, 178))

    d.line([(margin, 1970), (W - margin, 1970)], fill=(255, 255, 255, 38), width=1)
    d.text((margin, 2010), f"{AUTHOR} · {TEAM}", font=_font("segoeuib.ttf", 34), fill=(245, 248, 250))
    d.text((margin, 2065), COURSE, font=_font("segoeui.ttf", 22), fill=(165, 184, 200))
    d.text((margin, 2110), "Octubre de 2026", font=_font("segoeui.ttf", 22), fill=(165, 184, 200))
    d.text((W - margin - 345, 2110), "Python · scikit-learn · CatBoost", font=_font("segoeui.ttf", 20),
           fill=(120, 145, 164))

    path = REPORTS / "cover.png"
    img.save(path, quality=95)
    return path


# ---------------------------------------------------------------------------
# Equations
# ---------------------------------------------------------------------------
EQUATIONS = {
    "tree": (r"\hat{y}(x)=\sum_{m=1}^{M} c_m\,\mathbf{1}\{x\in R_m\}", 12.5),
    "variance": (r"\Delta I(s)=I(R)-\frac{|R_L|}{|R|}\,I(R_L)-\frac{|R_R|}{|R|}\,I(R_R)", 14.0),
    "extra": (r"\hat{y}_{\mathrm{ET}}(x)=\frac{1}{M}\sum_{m=1}^{M}T_m(x)", 9.0),
    "alpha": (r"\varepsilon_t=\sum_i w_i^{(t)}\,\mathbf{1}[h_t(x_i)\neq y_i],\qquad \alpha_t=\frac{1}{2}\ln\frac{1-\varepsilon_t}{\varepsilon_t}", 16.0),
    "update": (r"w_i^{(t+1)}=\frac{w_i^{(t)}\exp(-\alpha_t y_i h_t(x_i))}{Z_t}", 10.0),
    "margin": (r"\hat{y}(x)=\mathrm{sign}\left(\sum_{t=1}^{T}\alpha_t h_t(x)\right)", 10.5),
    "additive": (r"F_m(x)=F_{m-1}(x)+\nu\,h_m(x)", 9.0),
    "softmax": (r"p_k(x)=\frac{\exp(F_k(x))}{\sum_{j=1}^{K}\exp(F_j(x))}", 9.5),
    "logloss": (r"\mathcal{L}=-\sum_{i=1}^{n}\sum_{k=1}^{K}y_{ik}\ln p_{ik}", 10.0),
    "newton": (r"w_j^{\star}=-\frac{\sum_{i\in R_j}g_i}{\sum_{i\in R_j}h_i+\lambda}", 8.5),
    "ordered": (r"\hat{x}_{i,k}=\frac{\sum_{j:\sigma(j)<\sigma(i)}\mathbf{1}[x_{j,k}=x_{i,k}]\,y_j+a\,p}{\sum_{j:\sigma(j)<\sigma(i)}\mathbf{1}[x_{j,k}=x_{i,k}]+a}", 16.0),
    "metrics": (r"\mathrm{RMSE}=\sqrt{\frac{1}{n}\sum_i(y_i-\hat{y}_i)^2},\qquad \mathrm{MAE}=\frac{1}{n}\sum_i|y_i-\hat{y}_i|", 16.0),
    "r2": (r"R^2=1-\frac{\sum_i(y_i-\hat{y}_i)^2}{\sum_i(y_i-\bar{y})^2}", 9.0),
    "f1": (r"F_1=2\cdot\frac{\mathrm{precision}\cdot\mathrm{recall}}{\mathrm{precision}+\mathrm{recall}}", 11.0),
}


def build_equations() -> None:
    for name, (tex, width_cm) in EQUATIONS.items():
        fig = plt.figure(figsize=(width_cm / 2.54, 0.9), dpi=300)
        fig.patch.set_facecolor("white")
        fig.text(0.5, 0.5, f"${tex}$", ha="center", va="center",
                 fontsize=27, color="#12202f")
        fig.savefig(EQ / f"{name}.png", dpi=300, bbox_inches="tight",
                    pad_inches=0.045, facecolor="white")
        plt.close(fig)


# ---------------------------------------------------------------------------
# DOCX helpers
# ---------------------------------------------------------------------------
def set_cell_bg(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=45, bottom=45, left=85, right=85) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for tag, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_table_borders(table, color="D8DEE7", size=4) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def set_col_widths(table, widths_cm) -> None:
    table.autofit = False
    for row in table.rows:
        for i, width in enumerate(widths_cm):
            row.cells[i].width = Cm(width)
    for i, width in enumerate(widths_cm):
        table.columns[i].width = Cm(width)


def shade_run(run, hex_color="EEF2F6") -> None:
    r_pr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    r_pr.append(shd)


TOKEN = re.compile(r"(\*\*.*?\*\*|`[^`]+`)")


def add_rich(paragraph, text: str, size: float | None = None,
             color: RGBColor | None = None, italic: bool = False) -> None:
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = FONT_CODE
            run.font.size = Pt((size or 10.5) - 1.2)
            shade_run(run)
        else:
            run = paragraph.add_run(part)
        run.font.name = FONT_BODY if not part.startswith("`") else FONT_CODE
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        run.italic = italic


def add_page_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(txt)
    run._r.append(fld_end)


def add_bottom_border(paragraph, color="CBD5E1", size=8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    specs = {
        "Heading 1": (17, NAVY, True, 20, 8),
        "Heading 2": (13, TEAL, True, 14, 6),
        "Heading 3": (11.5, NAVY, True, 10, 4),
    }
    for name, (size, color, bold, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = FONT_DISPLAY
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.font.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_DISPLAY)

    cap = doc.styles["Caption"]
    cap.font.name = FONT_BODY
    cap.font.size = Pt(8.5)
    cap.font.italic = True
    cap.font.color.rgb = MUTED
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)


def add_section_number(doc: Document, number: str, title: str, first: bool = False) -> None:
    p = doc.add_paragraph(style="Heading 1")
    if not first:
        p.paragraph_format.page_break_before = True
    p.add_run(f"{number}. {title}")
    add_bottom_border(p, color="D8DEE7", size=6)


def add_h2(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="Heading 2")


def add_h3(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="Heading 3")


def body(doc: Document, text: str, space_after: float = 5, size: float = 10.5,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.alignment = align
    add_rich(p, text, size=size)


def bullets(doc: Document, items, size: float = 10) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, item, size=size)


def numbered(doc: Document, items, size: float = 10) -> None:
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(2)
        add_rich(p, f"**{i}.** {item}", size=size)


def callout(doc: Document, label: str, text: str, fill="F1F7F6", border="0F766E") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=border, size=8)
    cell = table.cell(0, 0)
    set_cell_bg(cell, fill)
    set_cell_margins(cell, top=110, bottom=110, left=160, right=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = TEAL
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_rich(p2, text, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers, rows, widths, font_size=8.5,
              header_bg="0A1928", zebra="F6F8FA", first_col_bold=False) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    keep_row_together(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_bg(cell, header_bg)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(str(text))
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = WHITE
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        keep_row_together(table.rows[-1])
        for ci, value in enumerate(row):
            cell = cells[ci]
            set_cell_margins(cell)
            if ri % 2 == 1:
                set_cell_bg(cell, zebra)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            txt = str(value)
            r = p.add_run(txt)
            r.font.size = Pt(font_size)
            if first_col_bold and ci == 0:
                r.bold = True
    set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def figure(doc: Document, filename: str, caption: str, width_cm=16.0) -> None:
    path = FIGURES / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph(style="Caption")
    m = re.match(r"^(Figura \d+\.)\s*(.*)$", caption)
    if m:
        r = cap.add_run(m.group(1) + " ")
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = INK
        add_rich(cap, m.group(2), size=8.5)
    else:
        add_rich(cap, caption, size=8.5, color=MUTED)


def equation(doc: Document, name: str, number: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(EQ / f"{name}.png"), width=Cm(EQUATIONS[name][1]))
    p.add_run("   ")
    r = p.add_run(f"({number})")
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "Actualice el índice en Word (F9) si los números de página no aparecen."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, txt, fld_end):
        run._r.append(el)


def code_block(doc: Document, code: str, size=8) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="D8DEE7", size=4)
    cell = table.cell(0, 0)
    set_cell_bg(cell, "0F1E2B")
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code)
    r.font.name = FONT_CODE
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("Árboles y Ensambles · Extra Trees, AdaBoost y CatBoost")
    r.font.name = FONT_BODY
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    hp.add_run("\t")
    r2 = hp.add_run(f"{TEAM} · {AUTHOR}")
    r2.font.name = FONT_BODY
    r2.font.size = Pt(8)
    r2.font.color.rgb = MUTED
    add_bottom_border(hp, color="D8DEE7", size=4)

    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run("Informe técnico · Octubre 2026")
    r.font.name = FONT_BODY
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    fp.add_run("\t")
    r = fp.add_run("Página ")
    r.font.name = FONT_BODY
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    add_page_field(fp, "PAGE")
    r = fp.add_run(" de ")
    r.font.name = FONT_BODY
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    add_page_field(fp, "NUMPAGES")


# ---------------------------------------------------------------------------
# Document body
# ---------------------------------------------------------------------------
def build_document() -> Path:
    cover = build_cover()
    build_equations()

    doc = Document()
    setup_styles(doc)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(0))
    sec.header_distance = Cm(0)
    sec.footer_distance = Cm(0)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(cover), width=Cm(21.0), height=Cm(29.6))

    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.page_width, sec2.page_height = Cm(21.0), Cm(29.7)
    sec2.left_margin = Cm(2.25)
    sec2.right_margin = Cm(2.25)
    sec2.top_margin = Cm(2.15)
    sec2.bottom_margin = Cm(1.95)
    sec2.header_distance = Cm(1.0)
    sec2.footer_distance = Cm(0.8)
    header_footer(sec2)

    # Executive summary -----------------------------------------------------
    add_section_number(doc, "0", "Resumen ejecutivo", first=True)
    body(doc,
         "Este informe documenta el diseño, el entrenamiento y la evaluación de tres ensambles de "
         "árboles aplicados a tres problemas reales con estructuras estadísticas distintas: "
         "pronóstico de ventas diarias en retail (Walmart M5), cribado de diabetes en salud pública "
         "(CDC BRFSS 2015) y predicción de abandono universitario (UCI Student Dropout). "
         "El objetivo no es declarar un ganador absoluto, sino mostrar **cuándo y por qué** cada "
         "familia de ensambles es la elección correcta, con evidencia cuantitativa y una simulación "
         "interactiva que permite auditar el comportamiento del modelo en el navegador.")
    body(doc,
         f"**Extra Trees** se evaluó sobre una muestra estratificada de 800 series de producto-tienda "
         f"(400 000 filas, 32 variables) con un split temporal de 355 200 / 22 400 / 22 400 filas. "
         f"El modelo final (366 árboles, profundidad 36, 70 % de variables por split) obtuvo "
         f"RMSE = {num(EXTRA['metrics']['rmse'])} unidades/día, MAE = {num(EXTRA['metrics']['mae'])}, "
         f"R² = {num(EXTRA['metrics']['r2'])} y WAPE = {num(EXTRA['metrics']['wape'], 1)} %. "
         f"Frente al baseline estacional `lag_7` (RMSE {num(EXTRA['metrics']['baseline']['seasonal_naive_lag7']['rmse'])}), "
         f"la reducción de error es de {pct(1 - EXTRA['metrics']['rmse'] / EXTRA['metrics']['baseline']['seasonal_naive_lag7']['rmse'], 1)}.")
    body(doc,
         f"**AdaBoost** se entrenó sobre 69 057 encuestas balanceadas de BRFSS (25 variables). "
         f"La búsqueda seleccionó 370 estimadores débiles de profundidad 3, con learning rate 0.10. "
         f"En test, ROC-AUC = {num(ADA['metrics']['roc_auc'])} y F1 = {num(ADA['metrics']['f1'])} con umbral 0.50. "
         f"El umbral operativo 0.38, elegido en una partición interna de entrenamiento, eleva el recall "
         f"a {num(ADA['metrics']['tuned']['recall'])} y el F1 a {num(ADA['metrics']['tuned']['f1'])}, "
         f"a costa de una precisión de {num(ADA['metrics']['tuned']['precision'])}. El Brier score es "
         f"{num(ADA['metrics']['brier'])}, lo que exige recalibración antes de comunicar probabilidades individuales.")
    body(doc,
         f"**CatBoost** se aplicó a 4 424 estudiantes con 39 variables (18 categóricas nativas). "
         f"El modelo final usó 172 árboles tras early stopping, profundidad 5 y pesos de clase balanceados. "
         f"En test alcanzó accuracy = {num(CAT['metrics']['accuracy'])}, balanced accuracy = {num(CAT['metrics']['balanced_accuracy'])}, "
         f"macro-F1 = {num(CAT['metrics']['f1_macro'])} y AUC-OvR = {num(CAT['metrics']['roc_auc_ovr'])}. "
         f"La clase Enrolled es la más difícil (F1 = {num(CAT['metrics']['per_class']['Enrolled']['f1'])}), "
         f"mientras Dropout y Graduate superan 0.74 de F1.")
    callout(doc, "Entregables",
            f"Repositorio con código y datos descargables: {REPO}. "
            f"Dashboard interactivo: {PAGES}. Informe en Word y PDF, figuras, tablas y modelos "
            f"exportados a JSON con verificación de paridad de inferencia (diferencia máxima "
            f"< 5 × 10⁻¹⁶).")
    body(doc, "**Palabras clave:** Extra Trees, AdaBoost, CatBoost, boosting, bagging, demanda intermitente, cribado, abandono universitario, calibración, interpretabilidad.", size=9.5)

    # TOC -------------------------------------------------------------------
    add_section_number(doc, "·", "Contenido")
    add_toc(doc)

    # 1 ---------------------------------------------------------------------
    add_section_number(doc, "1", "Introducción")
    add_h2(doc, "1.1 Contexto y motivación")
    body(doc,
         "Los árboles de decisión son atractivos porque aproximan funciones no lineales sin imponer una forma "
         "paramétrica, toleran escalas heterogéneas y producen reglas auditables. Su debilidad es la varianza: "
         "un árbol profundo cambia mucho ante pequeñas perturbaciones de los datos. Los ensambles atacan ese "
         "problema por dos caminos. El **bagging** promedia muchos árboles entrenados sobre muestras o "
         "perturbaciones independientes y reduce la varianza. El **boosting** construye árboles de forma "
         "secuencial, cada uno corrigiendo los errores del anterior, y reduce el sesgo. Extra Trees, AdaBoost "
         "y CatBoost son tres expresiones maduras de esas ideas, con supuestos y costos distintos.")
    body(doc,
         "La elección de los tres casos no es accidental. M5 representa un problema de regresión con demanda "
         "intermitente, muchos ceros, estacionalidad semanal y precios; BRFSS es un problema binario con "
         "variables clínicas y de estilo de vida, donde el costo de un falso negativo es alto; UCI Dropout es "
         "un problema multiclase con categorías nominales de alta cardinalidad (curso, modo de admisión, "
         "cualificación de los padres) y una clase intermedia ambigua. Cada estructura exige una estrategia "
         "de validación y una métrica diferente.")
    add_h2(doc, "1.2 Objetivos")
    bullets(doc, [
        "**Objetivo general:** implementar, ajustar y comparar Extra Trees, AdaBoost y CatBoost en tres casos de uso reales, con un pipeline reproducible y una interfaz web que permita reproducir el entrenamiento y simular inferencias.",
        "**Objetivos específicos:** (i) documentar la formulación matemática y los supuestos de cada ensamble; (ii) construir splits sin fuga de información y justificar las métricas; (iii) ajustar hiperparámetros con validación cruzada y reportar intervalos de confianza; (iv) analizar críticamente el valor y los límites de cada modelo en su dominio; (v) publicar código, artefactos y una simulación interactiva.",
    ])
    add_h2(doc, "1.3 Preguntas de investigación")
    numbered(doc, [
        "¿Qué compromiso entre sesgo, varianza e interpretabilidad ofrece cada ensamble en un régimen de datos distinto?",
        "¿Cómo cambia la decisión operativa cuando se elige un umbral o una métrica en lugar de otra?",
        "¿Qué evidencia se necesita para pasar de un buen resultado offline a una recomendación responsable en retail, salud o educación?",
    ])
    add_h2(doc, "1.4 Aportes y organización")
    body(doc,
         "El trabajo aporta (i) un pipeline reproducible con semillas fijas y separación estricta de train, "
         "validación y test; (ii) una comparación cuantitativa con métricas adecuadas a cada tarea; "
         "(iii) una exportación de los modelos a JavaScript con verificación numérica de paridad; y "
         "(iv) un análisis crítico de aplicaciones que explicita riesgos, sesgos y condiciones de uso. "
         "El resto del informe se organiza así: la Sección 2 presenta los fundamentos matemáticos; la "
         "Sección 3 describe datos y metodología; la Sección 4 reporta resultados; la Sección 5 analiza "
         "aplicaciones y límites; la Sección 6 concluye; los anexos contienen grillas, comandos y estructura del repositorio.")

    # 2 ---------------------------------------------------------------------
    add_section_number(doc, "2", "Fundamentos matemáticos")
    add_h2(doc, "2.1 Árboles de decisión y particiones recursivas")
    body(doc,
         "Un árbol de regresión divide el espacio de entrada en regiones disjuntas R₁, …, R_M mediante "
         "reglas binarias del tipo x_j ≤ s. En cada región, la predicción es una constante c_m, normalmente "
         "la media de las observaciones que caen en ella:")
    equation(doc, "tree", 1)
    body(doc,
         "La partición se elige de forma greedy: en cada nodo se evalúan candidatos (variable, umbral) y se "
         "maximiza la reducción de impureza. Para regresión, la impureza es la varianza o el error cuadrático "
         "medio; para clasificación, el índice Gini o la entropía. La ganancia de un split s es:")
    equation(doc, "variance", 2)
    body(doc,
         "donde R_L y R_R son los nodos hijos. El árbol crece hasta una profundidad máxima, un tamaño mínimo "
         "de hoja o una ganancia mínima. Sin poda, un árbol profundo tiene bajo sesgo y alta varianza: es un "
         "estimador inestable. Los ensambles estabilizan esa estimación promediando o corrigiendo árboles.")
    body(doc,
         "La descomposición sesgo-varianza del error cuadrático esperado es "
         "E[(y − f̂(x))²] = σ² + (f(x) − E[f̂(x)])² + Var(f̂(x)). El bagging reduce el tercer término al "
         "promediar modelos aproximadamente independientes; el boosting reduce el segundo al ajustar "
         "residuales de forma secuencial. Ninguno elimina el ruido irreducible σ² ni corrige por sí solo "
         "un sesgo de selección en los datos.")

    add_h2(doc, "2.2 Extra Trees: aleatoriedad extrema y promediado")
    body(doc,
         "Extra Trees (Extremely Randomized Trees) es un ensamble de tipo bagging con dos diferencias "
         "respecto de Random Forest. Primero, no remuestrea con reemplazo: cada árbol se entrena con el "
         "conjunto completo (o una muestra sin reemplazo, según la implementación). Segundo, en cada nodo "
         "no busca el mejor umbral entre todos los candidatos: sortea varios umbrales por variable y elige "
         "el mejor del sorteo. Esa aleatoriedad adicional reduce la correlación entre árboles y suele "
         "disminuir la varianza, a costa de un ligero aumento del sesgo.")
    body(doc, "La predicción de regresión es el promedio de M árboles independientes:")
    equation(doc, "extra", 3)
    body(doc,
         "La aleatoriedad de los umbrales actúa como regularización: impide que todos los árboles exploten "
         "el mismo punto de corte óptimo y hace que el ensamble sea menos sensible a ruido y outliers. "
         "En scikit-learn, `ExtraTreesRegressor` expone `max_features`, `min_samples_leaf`, "
         "`min_samples_split`, `max_depth` y `n_estimators`. La complejidad de entrenamiento es "
         "O(M · n log n) en el caso balanceado y el método es paralelizable por árbol.")
    callout(doc, "Por qué encaja en M5",
            "La demanda diaria tiene una masa grande de ceros, picos y una estacionalidad semanal fuerte. "
            "Un modelo lineal subestima la no linealidad; un árbol único sobreajusta. Extra Trees promedia "
            "muchos árboles poco correlacionados y tolera los ceros sin transformaciones especiales. "
            "El precio es que suaviza picos extremos y no extrapola tendencias fuera del rango observado.")

    add_h2(doc, "2.3 AdaBoost: boosting adaptativo y pérdida exponencial")
    body(doc,
         "AdaBoost combina clasificadores débiles h_t (típicamente árboles poco profundos) con pesos α_t. "
         "En cada ronda t, cada observación i tiene un peso w_i^{(t)}. El clasificador débil minimiza el "
         "error ponderado ε_t; su peso en el ensamble crece cuando ε_t es pequeño:")
    equation(doc, "alpha", 4)
    body(doc,
         "Después de entrenar h_t, los pesos se actualizan para castigar los errores y se normalizan por "
         "Z_t para que sumen uno:")
    equation(doc, "update", 5)
    body(doc, "La predicción final es el signo del voto ponderado:")
    equation(doc, "margin", 6)
    body(doc,
         "AdaBoost puede verse como un descenso por coordenadas en el espacio de funciones que minimiza la "
         "pérdida exponencial L(y, f) = exp(−y f(x)). Si cada clasificador débil tiene margen γ, es decir, "
         "ε_t ≤ ½ − γ, el error de entrenamiento está acotado por exp(−2Tγ²). Esa cota explica por qué "
         "añadir rondas puede seguir mejorando el ajuste, pero también por qué el método es sensible a "
         "etiquetas ruidosas: un ejemplo mal etiquetado recibe pesos crecientes y puede dominar una ronda.")
    body(doc,
         "En scikit-learn, `AdaBoostClassifier` con `algorithm='SAMME'` (el usado aquí) combina los votos "
         "y convierte la función de decisión en probabilidad mediante una sigmoide calibrada de forma "
         "aproximada. En la exportación a JavaScript se replica exactamente esa transformación: "
         "p = σ(2 · Σ α_t sign(h_t) / Σ α_t). Por eso las probabilidades deben interpretarse como scores "
         "de riesgo, no como probabilidades clínicas calibradas.")
    callout(doc, "Elección del umbral",
            "El umbral 0.50 maximiza exactitud, pero en cribado el costo de un falso negativo es mayor. "
            "Se eligió 0.38 sobre una partición interna del 15 % del entrenamiento, nunca sobre test. "
            "El resultado es un clasificador con recall 0.912 y precisión 0.670: apto para priorizar "
            "derivaciones, no para diagnosticar.")

    add_h2(doc, "2.4 CatBoost: boosting con categóricas nativas")
    body(doc,
         "CatBoost es un gradient boosting que construye un modelo aditivo. Para clasificación multiclase "
         "con K clases, la predicción se obtiene con una softmax sobre las funciones F_k:")
    equation(doc, "additive", 7)
    equation(doc, "softmax", 8)
    body(doc, "La pérdida es la entropía cruzada multiclase:")
    equation(doc, "logloss", 9)
    body(doc,
         "En cada iteración se ajusta un árbol a los gradientes g_i = p_i − y_i y a los hessianos "
         "h_i = p_i(1 − p_i). Cada hoja recibe el valor Newton que minimiza la aproximación cuadrática local:")
    equation(doc, "newton", 10)
    body(doc,
         "La innovación de CatBoost es el tratamiento de variables categóricas. En lugar de un one-hot "
         "de alta dimensión, reemplaza cada categoría por una estadística objetivo calculada con una "
         "permutación aleatoria y solo con ejemplos anteriores en esa permutación. Para la categoría k:")
    equation(doc, "ordered", 11)
    body(doc,
         "donde a es un parámetro de suavizado y p es el prior. Esa restricción evita el target leakage "
         "del target encoding clásico y permite usar categorías de alta cardinalidad sin explosión de "
         "columnas. Los árboles son *oblivious*: el mismo split se aplica en todo un nivel, lo que reduce "
         "el sobreajuste, acelera la inferencia y produce importancias estables. El boosting se controla "
         "con shrinkage ν, regularización L2 λ, profundidad y early stopping.")
    callout(doc, "Por qué encaja en UCI",
            "El dataset tiene 18 variables categóricas, varias con decenas de niveles (curso, modo de "
            "admisión, cualificación de los padres). Un one-hot clásico diluye la señal y crea columnas "
            "casi vacías. CatBoost aprende una representación ordenada de cada categoría y comparte "
            "estadísticas entre niveles, lo que mejora la generalización con solo 4 424 filas.")

    add_h2(doc, "2.5 Métricas y validación")
    body(doc,
         "En regresión se reportan RMSE, MAE, R² y WAPE. RMSE penaliza más los errores grandes y tiene "
         "las unidades de la variable objetivo; MAE es robusto y fácil de comunicar; R² compara con la "
         "media; WAPE (suma de errores absolutos / suma de valores reales) es útil cuando hay ceros y "
         "MAPE se vuelve inestable.")
    equation(doc, "metrics", 12)
    equation(doc, "r2", 13)
    body(doc,
         "En clasificación binaria se reportan accuracy, precisión, recall (sensibilidad), especificidad, "
         "F1, ROC-AUC, average precision, Brier score y log-loss. El F1 combina precisión y recall:")
    equation(doc, "f1", 14)
    body(doc,
         "En multiclase se usa macro-F1 (promedio no ponderado por clase), balanced accuracy y AUC-OvR. "
         "Macro-F1 es exigente con las clases minoritarias: una clase con mal desempeño arrastra el "
         "promedio. Los intervalos de confianza se estiman con bootstrap percentil (2 000 remuestreos "
         "para regresión y 500–1 000 para clasificación), y el umbral de AdaBoost se selecciona en una "
         "partición interna, no en test. La validación de Extra Trees es temporal (últimos 28 días de "
         "validación y 28 de test por serie); la de AdaBoost y CatBoost es estratificada 80/20.")

    # 3 ---------------------------------------------------------------------
    add_section_number(doc, "3", "Datos y metodología")
    add_h2(doc, "3.1 Conjuntos de datos")
    body(doc,
         "La Tabla 1 resume los tres datasets. Los tamaños son los efectivamente usados después de "
         "limpieza; no se reportan cifras nominales de los repositorios originales.")
    add_table(doc,
              ["Caso", "Fuente", "Tarea", "Filas", "Variables", "Clases / target"],
              [
                  ["Extra Trees", "Walmart M5", "Regresión", "400 000", "32", "ventas diarias (0–141)"],
                  ["AdaBoost", "CDC BRFSS 2015", "Binaria", "69 057", "25", "diabetes / no diabetes"],
                  ["CatBoost", "UCI Student Dropout", "Multiclase", "4 424", "39", "Dropout / Enrolled / Graduate"],
              ],
              [2.5, 3.7, 2.4, 2.1, 2.0, 4.0], font_size=8.5)
    body(doc, "**Tabla 1.** Resumen de los datasets y de la tarea supervisada.", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    add_h3(doc, "Walmart M5")
    body(doc,
         "M5 contiene 30 490 series de producto-tienda y un calendario de 1 969 días; la extracción "
         "de ventas usada abarca 1 913 días (2011-01-29 a 2016-04-24). Para un experimento docente "
         "reproducible se tomó una muestra estratificada de 80 series por tienda (800 series, 10 tiendas) "
         "y la ventana de 500 días comprendida entre 2014-12-12 y 2016-04-24. El target es la venta diaria en unidades. La muestra conserva la estructura "
         "jerárquica (tienda, departamento, categoría, estado), el calendario SNAP y los eventos, y el "
         "precio de venta semanal.")
    bullets(doc, [
        "**Split temporal:** 355 200 filas de entrenamiento, 22 400 de validación (28 días) y 22 400 de test (28 días). Nunca se mezclan fechas.",
        "**Features (32):** identificadores categóricos, calendario (día de semana, mes, día del año, fin de semana, SNAP, evento), precio y variaciones, lags 1/7/14/28/364, medias y desviaciones móviles 7/28/90, máximo móvil, ratio de ceros y estadísticas históricas item/tienda/departamento.",
        "**Fuga de información:** las estadísticas históricas se estiman solo con el período anterior a la ventana; los lags usan la historia completa previa y se desplazan por serie.",
        "**Distribución:** 59.75 % de ceros, media 1.31 y desviación 3.18 unidades/día; el máximo es 141. Es una distribución cero-inflada y muy asimétrica.",
    ])

    add_h3(doc, "CDC BRFSS 2015")
    body(doc,
         "La versión balanceada del Indicador de Diabetes de BRFSS contiene 70 692 respuestas con 21 "
         "variables demográficas, clínicas y de estilo de vida. Tras eliminar 1 635 duplicados quedan "
         "69 057 filas, con una tasa de positivos de 50.8 %. El target `Diabetes_binary` indica diagnóstico "
         "de diabetes o prediabetes. Se añadieron cuatro variables derivadas: comorbilidad (suma de "
         "hipertensión, colesterol alto, derrame y enfermedad cardíaca), hábitos saludables, categoría de "
         "IMC y días de salud mental+física. El split es estratificado 80/20 (13 812 filas de test).")
    bullets(doc, [
        "**Por qué no se usa el dataset completo de 253 680:** la versión balanceada facilita la comparación y evita que un clasificador trivial prediga la clase mayoritaria; se documenta como decisión de diseño.",
        "**Valores:** todas las variables son numéricas o binarias; no hay faltantes explícitos. El escalado se ajusta dentro del pipeline, solo con train.",
        "**Riesgo:** la encuesta es autodeclarada y el balanceo artificial no representa la prevalencia poblacional. Las probabilidades no deben interpretarse como riesgo clínico absoluto.",
    ])

    add_h3(doc, "UCI Student Dropout")
    body(doc,
         "El dataset contiene 4 424 estudiantes y 36 variables originales (18 categóricas y 18 numéricas) "
         "sobre estado civil, modo de admisión, curso, cualificaciones y ocupaciones de los padres, edad, "
         "notas de admisión, unidades curriculares aprobadas/matriculadas y variables macroeconómicas. "
         "Se añadieron tres ratios derivados: aprobación en primer semestre, aprobación en segundo "
         "semestre y cambio de nota entre semestres, hasta 39 variables. Las clases son Dropout (1 421), "
         "Enrolled (794) y Graduate (2 209). El split estratificado 80/20 deja 885 estudiantes en test.")
    bullets(doc, [
        "**Categóricas nativas:** CatBoost recibe los 18 campos como categóricos, sin one-hot manual; el modelo demo del navegador usa one-hot para poder exportarse a JSON de forma compacta.",
        "**Clase ambigua:** Enrolled es una categoría intermedia entre Dropout y Graduate; su F1 cercano a 0.50 no es un fallo aislado, sino una consecuencia de la definición del target.",
        "**Privacidad:** el dataset es público y anonimizado, pero cualquier despliegue real debe evaluar reidentificación y consentimiento.",
    ])

    add_h2(doc, "3.2 Protocolo de entrenamiento y búsqueda")
    body(doc,
         "Los tres experimentos usan semilla 42 y un pipeline de scikit-learn. La búsqueda de "
         "hiperparámetros se hizo con `RandomizedSearchCV` y validación cruzada. En Extra Trees, para "
         "contener el costo, se muestrearon 120 000 filas de entrenamiento y se usó una única partición "
         "explícita train/validación (sin fuga: el índice de validación queda fuera del de entrenamiento). "
         "En AdaBoost y CatBoost se usó `StratifiedKFold` de 3 folds. La Tabla 2 resume los espacios y "
         "el mejor punto encontrado.")
    add_table(doc,
              ["Modelo", "Búsqueda", "Mejor configuración", "Métrica CV"],
              [
                  ["Extra Trees", "12 configuraciones, 1 split temporal",
                   compact(EXTRA["search"][0]["config"]) if EXTRA.get("search") else "—",
                   f"RMSE {num(EXTRA['cv_rmse'])}"],
                  ["AdaBoost", "14 configuraciones, 3-fold estratificado",
                   compact(" | ".join(f"{k}={v}" for k, v in ADA["best_params"].items())),
                   f"AUC {num(ADA['cv_auc'])}"],
                  ["CatBoost", "14 configuraciones, 3-fold estratificado",
                   compact(" | ".join(f"{k}={v}" for k, v in CAT["best_params"].items())),
                   f"macro-F1 {num(CAT['cv_f1_macro'])}"],
              ],
              [2.5, 4.0, 7.1, 2.9], font_size=8.3)
    body(doc, "**Tabla 2.** Espacios de búsqueda y mejor configuración. Los valores completos están en los CSV de resultados y en el Anexo A.", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    body(doc,
         "La selección no se hizo por el mejor score de test, sino por validación cruzada. El test se "
         "usó una sola vez, después de fijar hiperparámetros y umbral. En Extra Trees, el modelo final "
         "se reentrenó con train+validación (práctica estándar para la entrega) y se evaluó en el bloque "
         "temporal de test; en AdaBoost y CatBoost, el modelo final se entrenó con el 80 % de train y se "
         "evaluó en el 20 % reservado. El umbral de AdaBoost se eligió en un 15 % interno del train, "
         "separado del test.")

    add_h2(doc, "3.3 Reproducibilidad y exportación web")
    body(doc,
         "Todo el pipeline se ejecuta con `python run_all.py`, que descarga los datos (si faltan), "
         "entrena los tres modelos, genera figuras y tablas y exporta los artefactos web. Las rutas y "
         "la semilla están centralizadas en `src/config.py`. Los modelos principales se guardan con "
         "joblib; el de Extra Trees pesa 2.9 GB y está excluido de Git, mientras que el modelo demo "
         "(44 árboles) se exporta a JSON de 7.4 MB para el navegador.")
    body(doc,
         "La exportación a JavaScript no es una caja negra: el script `src/export_web.py` implementa "
         "la misma lógica de inferencia en Python y en el JSON exportado, y compara predicciones en "
         "muestras de test. Las diferencias máximas absolutas son 4.4 × 10⁻¹⁶ para Extra Trees, 0.0 "
         "para AdaBoost y 5.6 × 10⁻¹⁷ para CatBoost, es decir, ruido de punto flotante. El dashboard "
         "permite reproducir las curvas de entrenamiento, ajustar perfiles y ver la contribución de "
         "cada árbol o variable sin depender de un servidor de modelos.")

    # 4 ---------------------------------------------------------------------
    add_section_number(doc, "4", "Resultados")
    body(doc,
         "Esta sección reporta los resultados en el conjunto de test y los contrasta con baselines "
         "sencillos. Las figuras se generaron a partir de los mismos artefactos que alimentan el "
         "dashboard; cualquier número del texto puede rastrearse hasta `artifacts/tables/*.json`.")

    add_h2(doc, "4.1 Extra Trees · Walmart M5")
    add_h3(doc, "Búsqueda y configuración")
    body(doc,
         f"La mejor configuración fue **{compact(EXTRA['search'][0]['config'])}**, con RMSE de validación "
         f"{num(EXTRA['search'][0]['rmse'])}. La diferencia con la segunda configuración es pequeña "
         f"(RMSE {num(EXTRA['search'][1]['rmse'])}), lo que sugiere una meseta: el modelo no es "
         "extremadamente sensible a los hiperparámetros dentro de la región explorada. El modelo final "
         "reentrenado con train+validación obtuvo el rendimiento de test de la Tabla 3.")
    add_table(doc,
              ["Métrica", "Extra Trees (test)", "Baseline lag_7", "Mejora"],
              [
                  ["RMSE (unidades/día)", num(EXTRA["metrics"]["rmse"]), num(EXTRA["metrics"]["baseline"]["seasonal_naive_lag7"]["rmse"]),
                   pct(1 - EXTRA["metrics"]["rmse"] / EXTRA["metrics"]["baseline"]["seasonal_naive_lag7"]["rmse"], 1)],
                  ["MAE (unidades/día)", num(EXTRA["metrics"]["mae"]), num(EXTRA["metrics"]["baseline"]["seasonal_naive_lag7"]["mae"]),
                   pct(1 - EXTRA["metrics"]["mae"] / EXTRA["metrics"]["baseline"]["seasonal_naive_lag7"]["mae"], 1)],
                  ["R²", num(EXTRA["metrics"]["r2"]), num(EXTRA["metrics"]["baseline"]["seasonal_naive_lag7"]["r2"]), "—"],
                  ["WAPE (%)", num(EXTRA["metrics"]["wape"], 1), num(EXTRA["metrics"]["baseline"]["seasonal_naive_lag7"]["wape"], 1), "—"],
                  ["Sesgo medio", num(EXTRA["metrics"]["bias"]), num(EXTRA["metrics"]["baseline"]["seasonal_naive_lag7"]["bias"]), "—"],
              ],
              [4.4, 4.1, 4.1, 3.5], font_size=8.8)
    body(doc, "**Tabla 3.** Métricas de Extra Trees en test frente al baseline estacional. La mejora se calcula sobre RMSE y MAE.", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    body(doc,
         f"El intervalo bootstrap del RMSE es [{num(EXTRA['metrics']['bootstrap']['rmse']['lo'])}, "
         f"{num(EXTRA['metrics']['bootstrap']['rmse']['hi'])}] y el de R² es "
         f"[{num(EXTRA['metrics']['bootstrap']['r2']['lo'])}, {num(EXTRA['metrics']['bootstrap']['r2']['hi'])}]. "
         "La mejora del 28 % frente al baseline es consistente, pero el WAPE de 70 % recuerda que la "
         "demanda intermitente es difícil: el modelo acierta la tendencia, no cada pico. El sesgo medio "
         "es −0.049 unidades/día, ligeramente negativo, sin señal de sobrepredicción sistemática.")
    callout(doc, "Lectura de negocio",
            "Con MAE ≈ 0.99 unidades/día, un gerente puede dimensionar reposición y stock de seguridad "
            "por producto. El modelo es más útil para ordenar prioridades y detectar series de alta "
            "rotación que para predecir un día exacto. El error relativo alto en series con muchos ceros "
            "obliga a complementar con reglas de negocio y revisión humana.")

    figure(doc, "m5_series_forecast.png", "Figura 1. Pronóstico de las tres series de mayor volumen en los 28 días de test. La línea oscura es la venta real y la ámbar la predicción.", 16.0)
    figure(doc, "m5_actual_vs_pred.png", "Figura 2. Real vs. predicho en test. La nube se concentra en ventas bajas, típicas de demanda intermitente; los puntos de alta demanda se acercan a la diagonal.", 14.5)
    figure(doc, "m5_learning_curve.png", "Figura 3. Curva de aprendizaje sobre el test completo (22 400 filas): el error cae rápido y se estabiliza alrededor de 2.0 unidades/día. La curva no usa el conjunto de validación que participó del entrenamiento final.", 14.5)
    figure(doc, "m5_importance.png", "Figura 4. Importancia por impureza. Los lags y las medias móviles dominan; el precio y el calendario aportan ajustes finos.", 14.5)
    figure(doc, "m5_pdp_2d.png", "Figura 5. Superficie de respuesta parcial (PDP) sobre lag_7 y la media móvil de 28 días. La demanda es casi monótona en ambas variables.", 13.5)

    add_h2(doc, "4.2 AdaBoost · CDC BRFSS")
    add_h3(doc, "Búsqueda y configuración")
    body(doc,
         f"La búsqueda eligió **n = {ADA['best_params'].get('n_estimators')}, learning rate = "
         f"{ADA['best_params'].get('learning_rate')}, profundidad = {ADA['best_params'].get('max_depth')} "
         f"y min_samples_leaf = {ADA['best_params'].get('min_samples_leaf')}**, con AUC de validación "
         f"cruzada {num(ADA['cv_auc'])}. El umbral operativo se fijó en {num(ADA['threshold'], 2)} sobre "
         "una partición interna del 15 % del entrenamiento. La Tabla 4 compara el umbral estándar 0.50 "
         "con el umbral operativo.")
    add_table(doc,
              ["Métrica", "Umbral 0.50", "Umbral 0.38", "Δ"],
              [
                  ["Accuracy", num(ADA["metrics"]["accuracy"]), num(ADA["metrics"]["tuned"]["accuracy"]), num(ADA["metrics"]["tuned"]["accuracy"] - ADA["metrics"]["accuracy"])],
                  ["Precisión", num(ADA["metrics"]["precision"]), num(ADA["metrics"]["tuned"]["precision"]), num(ADA["metrics"]["tuned"]["precision"] - ADA["metrics"]["precision"])],
                  ["Recall", num(ADA["metrics"]["recall"]), num(ADA["metrics"]["tuned"]["recall"]), num(ADA["metrics"]["tuned"]["recall"] - ADA["metrics"]["recall"])],
                  ["F1", num(ADA["metrics"]["f1"]), num(ADA["metrics"]["tuned"]["f1"]), num(ADA["metrics"]["tuned"]["f1"] - ADA["metrics"]["f1"])],
                  ["Especificidad", num(ADA["metrics"]["specificity"]), num(ADA["metrics"]["tuned"]["specificity"]), num(ADA["metrics"]["tuned"]["specificity"] - ADA["metrics"]["specificity"])],
                  ["Brier", num(ADA["metrics"]["brier"]), num(ADA["metrics"]["tuned"]["brier"]), "0.000"],
              ],
              [4.0, 4.2, 4.2, 3.4], font_size=8.8)
    body(doc, "**Tabla 4.** Compromiso entre umbrales en test. Bajar el umbral prioriza recall (no dejar casos sin detectar) a costa de precisión y especificidad.", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    body(doc,
         f"El AUC es {num(ADA['metrics']['roc_auc'])} y el average precision "
         f"{num(ADA['metrics']['average_precision'])}; el intervalo bootstrap del AUC es "
         f"[{num(ADA['metrics']['bootstrap']['roc_auc']['lo'])}, {num(ADA['metrics']['bootstrap']['roc_auc']['hi'])}] "
         f"y el del F1 es [{num(ADA['metrics']['bootstrap']['f1']['lo'])}, {num(ADA['metrics']['bootstrap']['f1']['hi'])}]. "
         "La ganancia de recall (0.780 → 0.912) es grande y consistente; la pérdida de precisión "
         "(0.737 → 0.670) es el precio de derivar más personas a confirmación. En cribado, ese "
         "intercambio suele ser aceptable si la prueba de confirmación es barata y el costo de un "
         "falso negativo es alto.")
    callout(doc, "Calibración",
            f"El Brier score es {num(ADA['metrics']['brier'])} y la curva de calibración muestra "
            "desviaciones en la zona alta. Las probabilidades sirven para ordenar riesgo, pero no para "
            "decir a una persona “tiene 63 % de probabilidad de diabetes”. Antes de un uso clínico se "
            "requiere recalibración (Platt o isotónica) y validación externa.")

    figure(doc, "cdc_roc.png", "Figura 6. Curva ROC en test. Un AUC de 0.821 indica buena separación entre personas con y sin diabetes.", 14.0)
    figure(doc, "cdc_confusion.png", "Figura 7. Matriz de confusión con umbral 0.38. Se prioriza no dejar casos positivos sin detectar; los falsos positivos se derivan a confirmación.", 12.0)
    figure(doc, "cdc_thresholds.png", "Figura 8. Barrido de umbral. Bajar de 0.50 a 0.38 eleva el recall de 0.780 a 0.912 a costa de precisión.", 15.0)
    figure(doc, "cdc_calibration.png", "Figura 9. Calibración de las probabilidades. La curva se aleja de la diagonal en los extremos; se recomienda recalibrar antes de comunicar riesgos individuales.", 13.5)
    figure(doc, "cdc_importance.png", "Figura 10. Importancia de variables. Salud general, edad, comorbilidad y BMI concentran la señal; los hábitos aportan información incremental.", 14.0)

    add_h2(doc, "4.3 CatBoost · UCI Student Dropout")
    add_h3(doc, "Búsqueda y configuración")
    body(doc,
         f"La mejor configuración fue **depth = {CAT['best_params'].get('depth')}, lr = "
         f"{CAT['best_params'].get('learning_rate')}, l2 = {CAT['best_params'].get('l2_leaf_reg')}, "
         f"iteraciones = {CAT['best_params'].get('iterations')}, random_strength = "
         f"{CAT['best_params'].get('random_strength')}, bagging_temperature = "
         f"{CAT['best_params'].get('bagging_temperature')}, auto_class_weights = Balanced**, con "
         f"macro-F1 de validación cruzada {num(CAT['cv_f1_macro'])}. El early stopping usó 172 árboles "
         "de los 350 configurados. La Tabla 5 reporta las métricas globales y por clase.")
    add_table(doc,
              ["Clase", "Precisión", "Recall", "F1", "Soporte"],
              [
                  [c, num(CAT["metrics"]["per_class"][c]["precision"]), num(CAT["metrics"]["per_class"][c]["recall"]), num(CAT["metrics"]["per_class"][c]["f1"]), str(CAT["metrics"]["per_class"][c]["support"])]
                  for c in ["Dropout", "Enrolled", "Graduate"]
              ],
              [4.0, 3.2, 3.2, 3.2, 2.2], font_size=9.0, first_col_bold=True)
    body(doc, "**Tabla 5.** Métricas por clase en test (885 estudiantes). Macro-F1 = 0.693; weighted-F1 = 0.744; accuracy = 0.731; balanced accuracy = 0.705.", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    body(doc,
         f"El AUC-OvR es {num(CAT['metrics']['roc_auc_ovr'])} y el log-loss "
         f"{num(CAT['metrics']['log_loss'])}. La clase Enrolled concentra la dificultad: su precisión "
         f"({num(CAT['metrics']['per_class']['Enrolled']['precision'])}) es baja porque comparte "
         "frontera con Dropout y Graduate. El modelo detecta bien Dropout (recall 0.694; F1 0.745) y "
         "Graduate (recall 0.792; F1 0.830), que son las decisiones de mayor consecuencia operativa. "
         "Los intervalos bootstrap son [0.703, 0.757] para accuracy y [0.664, 0.722] para macro-F1.")
    callout(doc, "Lectura de negocio",
            "El modelo no debe etiquetar a un estudiante como “abandono” de forma automática. Su valor "
            "es priorizar tutorías y revisar casos con un equipo humano. La clase Enrolled es una señal "
            "de que el sistema debe mirar trayectorias, no solo un snapshot: un estudiante en curso "
            "puede graduarse o abandonar con cambios pequeños en las variables.")

    figure(doc, "uci_confusion.png", "Figura 11. Matriz de confusión 3×3. La clase Enrolled es la más difícil: comparte frontera con Dropout y Graduate.", 12.0)
    figure(doc, "uci_per_class.png", "Figura 12. Precisión, recall y F1 por clase. Dropout y Graduate superan 0.74 de F1; Enrolled queda en 0.50.", 15.0)
    figure(doc, "uci_roc.png", "Figura 13. Curvas ROC one-vs-rest. Graduate es la clase más separable; Enrolled, la más ambigua.", 14.5)
    figure(doc, "uci_importance.png", "Figura 14. Importancia de variables. Las unidades aprobadas del segundo semestre, la matrícula al día y el rendimiento previo dominan la predicción.", 14.5)
    figure(doc, "uci_calibration.png", "Figura 15. Calibración de la probabilidad de Dropout. Es útil para priorizar tutorías, no para decisiones automáticas.", 13.0)

    add_h2(doc, "4.4 Comparación entre los tres ensambles")
    body(doc,
         "La Tabla 6 resume el desempeño y el costo. La comparación debe leerse con cautela: las tareas, "
         "los tamaños y las métricas no son intercambiables. Un RMSE de 2.0 unidades/día no es "
         "“mejor” que un AUC de 0.82; son respuestas a preguntas distintas.")
    add_table(doc,
              ["Modelo", "Tarea", "Test", "Métrica principal", "Métrica secundaria", "Modelo"],
              [
                  ["Extra Trees", "Regresión", "22 400", f"RMSE {num(EXTRA['metrics']['rmse'])}", f"R² {num(EXTRA['metrics']['r2'])}", "2.9 GB"],
                  ["AdaBoost", "Binaria", "13 812", f"AUC {num(ADA['metrics']['roc_auc'])}", f"F1 {num(ADA['metrics']['f1'])}", "0.57 MB"],
                  ["CatBoost", "Multiclase", "885", f"macro-F1 {num(CAT['metrics']['f1_macro'])}", f"AUC-OvR {num(CAT['metrics']['roc_auc_ovr'])}", "0.28 MB"],
              ],
              [2.8, 2.5, 2.4, 4.2, 4.0, 2.4], font_size=8.6)
    body(doc, "**Tabla 6.** Comparación de desempeño y tamaño de modelo. Los tamaños son los joblib principales; los modelos demo del navegador son más pequeños.", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    body(doc,
         "En términos de esfuerzo de ajuste, CatBoost fue el más costoso (búsqueda y entrenamiento "
         "completos en ~16 minutos), seguido de Extra Trees (~7 minutos) y AdaBoost (~1.7 minutos). "
         "En inferencia interactiva, los tres responden en milisegundos; la diferencia práctica está "
         "en el tamaño y en la necesidad de preprocesamiento. AdaBoost y CatBoost ofrecen probabilidades "
         "o scores; Extra Trees entrega una predicción continua y una banda de incertidumbre entre "
         "árboles. La Figura 16 evidencia además que la búsqueda de Extra Trees se estabiliza en una "
         "meseta de RMSE, por lo que el modelo no exige un ajuste fino para rendir de forma competitiva.")
    figure(doc, "m5_search.png", "Figura 16. Las 12 configuraciones de RandomizedSearchCV para Extra Trees. La mejor combina 366 árboles, profundidad 36 y 70 % de variables por split.", 14.5)
    body(doc,
         "La lectura integrada es que **no hay un ganador único**. Extra Trees es la opción natural para "
         "regresión tabular de alta dimensionalidad con ruido; AdaBoost es un baseline fuerte, barato "
         "e interpretable para clasificación binaria; CatBoost es la mejor opción cuando hay categóricas "
         "de alta cardinalidad y se necesita un modelo compacto y preciso. La sección 5 traduce esa "
         "conclusión a criterios de aplicación.")

    # 5 ---------------------------------------------------------------------
    add_section_number(doc, "5", "Aplicaciones y análisis crítico")
    body(doc,
         "Un modelo útil no es solo el que obtiene una métrica alta, sino el que mejora una decisión "
         "sin trasladar riesgos invisibles al usuario. Esta sección analiza cada caso desde esa "
         "perspectiva: qué decisión habilita, qué supuestos exige, qué puede fallar y qué controles "
         "humanos y técnicos necesita.")

    add_h2(doc, "5.1 Retail: inventario con demanda intermitente")
    body(doc,
         "En M5, la decisión es cuánto reponer y cómo priorizar el surtido. El modelo entrega una "
         "predicción diaria por producto-tienda; el MAE de ~1 unidad/día es accionable a nivel de "
         "categoría, pero no debe usarse como orden de compra automática. La demanda tiene 59.75 % de "
         "ceros: un error absoluto de 1 unidad puede ser enorme en una serie que vende 0.2 unidades/día "
         "y pequeño en una que vende 20. Por eso se reporta WAPE y se recomienda segmentar por volumen.")
    bullets(doc, [
        "**Valor:** priorización de reposición, detección de series de alta rotación y planificación de seguridad de stock. El modelo reduce el RMSE un 28 % frente a lag_7.",
        "**Límite:** no extrapola tendencias ni promociones fuera del rango histórico; suaviza picos. La muestra de 800 series no cubre las 30 490 del dataset completo.",
        "**Control:** comparar contra un baseline estacional, monitorear deriva por tienda/departamento y mantener un override humano para eventos y quiebres de stock.",
        "**Ética/operación:** una mala predicción puede generar faltantes o exceso de inventario; el costo asimétrico debe incorporarse en la función de decisión, no solo en la métrica.",
    ])
    body(doc,
         "La implementación razonable es un **sistema de recomendación**, no un piloto automático: el "
         "modelo sugiere, el gerente ajusta y el sistema registra el resultado para reentrenar. La "
         "validación debe ser temporal y por tienda; una validación aleatoria mezclaría días futuros "
         "con pasados y sobreestimaría el desempeño.")

    add_h2(doc, "5.2 Salud pública: cribado con recall alto")
    body(doc,
         "En BRFSS, la decisión es a quién invitar a una confirmación diagnóstica. El costo de un falso "
         "negativo es alto; el de un falso positivo es una prueba adicional y ansiedad. Por eso el "
         "umbral 0.38 es una decisión de política, no una propiedad del modelo: aumenta el recall a "
         "0.912 y reduce la precisión a 0.670. El Brier de 0.178 y la curva de calibración muestran que "
         "las probabilidades no son clínicamente exactas; el modelo debe usarse como ranking de riesgo.")
    bullets(doc, [
        "**Valor:** cribado poblacional de bajo costo, con un modelo ligero (0.57 MB) y entrenamiento en segundos; permite priorizar recursos limitados.",
        "**Límite:** la encuesta es autodeclarada; el balanceo 50/50 no representa la prevalencia real; el modelo no incorpora historia clínica ni biomarcadores.",
        "**Equidad:** auditar recall y precisión por edad, sexo, ingreso y educación. Un modelo con buen promedio puede fallar en subgrupos pequeños.",
        "**Gobernanza:** no es diagnóstico; requiere confirmación clínica, explicación al paciente y revisión de sesgos antes de cualquier despliegue.",
    ])

    add_h2(doc, "5.3 Educación: alerta temprana de abandono")
    body(doc,
         "En UCI, la decisión es qué estudiantes requieren acompañamiento. CatBoost detecta bien "
         "Dropout y Graduate, pero confunde Enrolled; además, el dataset es un snapshot con variables "
         "de dos semestres, no una secuencia temporal completa. Un despliegue real necesita validación "
         "temporal (predecir cohortes futuras), privacidad y un protocolo de intervención.")
    bullets(doc, [
        "**Valor:** prioriza tutorías y permite asignar recursos a estudiantes con mayor riesgo; la importancia de variables (aprobadas en 2.º semestre, matrícula al día) es accionable.",
        "**Límite:** 4 424 filas y 885 en test; la clase Enrolled es ambigua; el modelo puede perpetuar sesgos si se entrena con decisiones pasadas.",
        "**Ética:** una etiqueta de “abandono” puede estigmatizar. Debe comunicarse como probabilidad de necesitar apoyo, con explicaciones y derecho a revisión.",
        "**Control:** validación por cohorte, monitoreo de deriva y auditoría de equidad por programa, género y condiciones socioeconómicas.",
    ])

    add_h2(doc, "5.4 Guía de selección")
    add_table(doc,
              ["Situación", "Modelo recomendado", "Por qué"],
              [
                  ["Regresión tabular, ruido, muchos outliers, necesidad de paralelizar", "Extra Trees", "Bagging extremo, bajo costo de tuning, robusto y rápido; no extrapola."],
                  ["Clasificación binaria, baseline fuerte, interpretabilidad por rondas", "AdaBoost", "Ligero, secuencial, buena curva ROC; sensible a ruido y calibración imperfecta."],
                  ["Categóricas de alta cardinalidad, multiclase, modelo compacto", "CatBoost", "Target statistics ordenadas, oblivious trees, buen desempeño con pocos datos."],
                  ["Necesidad de probabilidades calibradas", "Recalibrar cualquiera", "Ninguno está calibrado por defecto; usar Platt/isotónica con validación externa."],
              ],
              [6.0, 3.6, 7.0], font_size=8.6, first_col_bold=False)
    body(doc, "**Tabla 7.** Guía práctica de selección. La decisión final depende de la métrica de negocio, el costo del error y los requisitos de explicabilidad.", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # 6 ---------------------------------------------------------------------
    add_section_number(doc, "6", "Conclusiones")
    body(doc,
         "Los tres ensambles confirman una idea central: la elección del modelo es inseparable de la "
         "estructura del problema y del costo de equivocarse. Extra Trees aprovecha la aleatoriedad "
         "para estabilizar una regresión difícil y reduce el error un 28 % frente a un baseline "
         "estacional, con un MAE cercano a una unidad/día. AdaBoost demuestra que un modelo pequeño y "
         "rápido puede separar casos de diabetes con AUC 0.821; su umbral operativo convierte una "
         "métrica en una política de cribado. CatBoost muestra que las categóricas nativas y el "
         "boosting ordenado permiten un modelo compacto y preciso en un dataset pequeño y multiclase, "
         "aunque la clase Enrolled revela el límite de la definición del target.")
    body(doc,
         "La comparación no produce un ganador absoluto. Produce un **mapa de decisión**: Extra Trees "
         "para regresión tabular ruidosa y paralelizable; AdaBoost para clasificación binaria con "
         "necesidad de recall y bajo costo; CatBoost para categóricas de alta cardinalidad y "
         "multiclase. En los tres casos, la validación honesta, la separación de test, el reporte de "
         "intervalos y la auditoría de sesgos son tan importantes como el algoritmo.")
    body(doc,
         "Quedan límites claros: la muestra de M5 no cubre todas las series; el umbral de AdaBoost se "
         "eligió en una partición pequeña; CatBoost tiene solo 885 estudiantes en test; las "
         "probabilidades no están calibradas para uso clínico. La reproducibilidad está garantizada "
         "por el repositorio, las semillas fijas, los scripts de descarga y la verificación de paridad "
         "entre Python y JavaScript.")
    add_h2(doc, "6.1 Trabajo futuro")
    bullets(doc, [
        "Validación temporal por cohorte y por tienda, con ventanas móviles y monitoreo de deriva.",
        "Recalibración de AdaBoost y CatBoost con Platt/isotónica y validación externa.",
        "Modelos jerárquicos para M5 (agregación por tienda/departamento) y funciones de pérdida asimétricas.",
        "Análisis de equidad por subgrupos y pruebas de estrés con datos sintéticos.",
        "Despliegue del dashboard como GitHub Pages y API de inferencia con límites de uso.",
    ])

    # References ------------------------------------------------------------
    add_section_number(doc, "7", "Referencias")
    refs = [
        "Breiman, L. (2001). Random forests. *Machine Learning, 45*(1), 5–32. https://doi.org/10.1023/A:1010933404324",
        "Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. *Proceedings of the 22nd ACM SIGKDD*, 785–794.",
        "Dorogush, A. V., Ershov, V., & Gulin, A. (2018). CatBoost: Gradient boosting with categorical features support. *arXiv:1810.11363*.",
        "Efron, B., & Tibshirani, R. J. (1993). *An introduction to the bootstrap*. Chapman & Hall.",
        "Freund, Y., & Schapire, R. E. (1997). A decision-theoretic generalization of on-line learning and an application to boosting. *Journal of Computer and System Sciences, 55*(1), 119–139.",
        "Friedman, J., Hastie, T., & Tibshirani, R. (2000). Additive logistic regression: A statistical view of boosting. *The Annals of Statistics, 28*(2), 337–407.",
        "Geurts, P., Ernst, D., & Wehenkel, L. (2006). Extremely randomized trees. *Machine Learning, 63*(1), 3–42.",
        "Hastie, T., Tibshirani, R., & Friedman, J. (2009). *The elements of statistical learning* (2nd ed.). Springer.",
        "James, G., Witten, D., Hastie, T., & Tibshirani, R. (2021). *An introduction to statistical learning* (2nd ed.). Springer.",
        "Makridakis, S., Spiliotis, E., & Assimakopoulos, V. (2022). M5 accuracy competition: Results, findings, and conclusions. *International Journal of Forecasting, 38*(4), 1346–1364.",
        "Pedregosa, F., et al. (2011). Scikit-learn: Machine learning in Python. *Journal of Machine Learning Research, 12*, 2825–2830.",
        "Prokhorenkova, L., Gusev, G., Vorobev, A., Dorogush, A. V., & Gulin, A. (2018). CatBoost: Unbiased boosting with categorical features. *Advances in Neural Information Processing Systems, 31*.",
        "Realinho, V., Machado, J., Baptista, L., & Martins, M. V. (2022). Predicting student dropout and academic success. *Data, 7*(11), 146.",
        "Saito, T., & Rehmsmeier, M. (2015). The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. *PLOS ONE, 10*(3), e0118432.",
        "Schapire, R. E., & Singer, Y. (1999). Improved boosting algorithms using confidence-rated predictions. *Machine Learning, 37*(3), 297–336.",
        "Centers for Disease Control and Prevention. (2015). *Behavioral Risk Factor Surveillance System survey data*. U.S. Department of Health and Human Services.",
        "UCI Machine Learning Repository. (2021). *Predict students' dropout and academic success* [Data set]. https://archive.ics.uci.edu/dataset/697/predict+students+dropout+and+academic+success",
        "Walmart. (2022). *M5 forecasting — accuracy* [Data set]. Kaggle. https://www.kaggle.com/competitions/m5-forecasting-accuracy",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, ref, size=9)

    # Annexes ---------------------------------------------------------------
    add_section_number(doc, "A", "Anexo A. Grillas de hiperparámetros")
    body(doc, "La Tabla A1 resume los espacios explorados. Los valores se implementaron con `scipy.stats` y listas discretas; el número de iteraciones de búsqueda se limitó para que el experimento fuera reproducible en un equipo personal.")
    add_table(doc,
              ["Modelo", "Espacio", "Iteraciones", "Validación"],
              [
                  ["Extra Trees", "n_estimators 150–400; max_depth 16–40; max_features 0.5/0.7/0.9; min_samples_leaf 2–24; min_samples_split 2–20; bootstrap True/False", "12", "1 split temporal (120k train / 22.4k val)"],
                  ["AdaBoost", "n_estimators 100–420; learning_rate 0.01–0.4; depth 1–4; min_samples_leaf 1/5/10/20", "14", "StratifiedKFold(3)"],
                  ["CatBoost", "iterations 250–650; depth 3–8; lr 0.02–0.10; l2 1–10; random_strength 0–2; bagging 0–1", "14", "StratifiedKFold(3)"],
              ],
              [2.6, 9.5, 1.9, 3.6], font_size=8.0)
    body(doc, "**Tabla A1.** Espacios de búsqueda. Los mejores puntos se muestran en la Tabla 2.", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    add_section_number(doc, "B", "Anexo B. Reproducibilidad y estructura")
    body(doc, "El repositorio se organiza así:")
    code_block(doc,
               "arboles-ensambles-equipo2/\n"
               "├── docs/                 # dashboard estático (GitHub Pages)\n"
               "│   ├── index.html\n"
               "│   ├── assets/{css,js}\n"
               "│   ├── data/dashboard.json\n"
               "│   └── models/*.json\n"
               "├── src/                  # pipeline Python\n"
               "│   ├── config.py\n"
               "│   ├── data_m5.py · data_cdc.py · data_uci.py\n"
               "│   ├── train_extra_trees.py\n"
               "│   ├── train_adaboost.py\n"
               "│   ├── train_catboost.py\n"
               "│   └── export_web.py\n"
               "├── artifacts/{figures,tables,models}\n"
               "├── reports/              # DOCX, PDF, ecuaciones\n"
               "├── scripts/download_data.py\n"
               "├── run_all.py\n"
               "└── requirements.txt",
               size=7.7)
    body(doc, "Comandos principales:")
    code_block(doc,
               "python -m venv .venv\n"
               ".venv\\Scripts\\activate\n"
               "pip install -r requirements.txt\n"
               "python scripts/download_data.py\n"
               "python run_all.py\n"
               "python -m src.make_report\n"
               "powershell -ExecutionPolicy Bypass -File scripts/convert_pdf.ps1",
               size=8)
    callout(doc, "Entorno de ejecución",
            "Experimentos ejecutados en Windows con Python 3.14, scikit-learn 1.9, CatBoost 1.2.10, "
            "NumPy 2.5 y pandas 3.0. Los tiempos de pared fueron "
            f"{num(EXTRA.get('runtime_seconds', 0), 1)} s para Extra Trees, "
            f"{num(ADA.get('runtime_seconds', 0), 1)} s para AdaBoost y "
            f"{num(CAT.get('runtime_seconds', 0), 1)} s para CatBoost, incluyendo búsqueda, "
            "entrenamiento, figuras y exportación. No se usó GPU; son un orden de magnitud, no un benchmark.")
    body(doc, "Los datos crudos no se versionan por tamaño; el script de descarga reconstruye `data/raw/`. Los modelos principales se regeneran con `run_all.py`; el joblib de Extra Trees (2.9 GB) queda fuera de Git y se conserva solo localmente. El dashboard funciona con cualquier servidor estático (`python -m http.server` en `docs/`).")

    add_section_number(doc, "C", "Anexo C. Declaración de uso responsable")
    body(doc, "Este trabajo es académico y reproducible. Los modelos no están validados para uso clínico, financiero o académico real. Cualquier despliegue debe someterse a revisión ética, validación externa, monitoreo de deriva y auditoría de equidad. Las decisiones que afectan a personas deben mantener revisión humana y mecanismos de apelación. Los datasets son públicos y anonimizados; aun así, deben respetarse las licencias y las condiciones de uso de cada fuente.")

    out = REPORTS / "Informe_Tecnico_Arboles_Ensambles.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    path = build_document()
    print("wrote", path, path.stat().st_size)
